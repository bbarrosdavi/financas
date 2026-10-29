"""Leitura dos arquivos do Bradesco: extrato da conta (.pdf, .csv ou .xlsx) e fatura do cartão em aberto (.pdf).

Extrato: tabela Data, Histórico, Docto., Crédito, Débito e Saldo. No PDF cada lançamento ocupa até
três linhas (histórico; data, valores e saldo; "REM:" ou "DES:" com a outra pessoa) e a data só
aparece no primeiro lançamento do dia. O saldo de cada linha prova a anterior: se a corrente de
saldos ou o total do PDF não fecham, o arquivo é recusado.

Fatura em aberto: o PDF que o aplicativo gera é só imagem (um print por página), então é lido por
OCR (tesseract com português). Data e valor de cada linha são relidos sozinhos, só com dígitos, e
cada cartão tem que fechar com o "Total para" impresso; se algum não fecha, nada entra.
"""

import io
import re
import shutil
import statistics
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from app.importadores import Leitura, Movimento, numero

CONTA, CARTAO = "bradesco_conta", "bradesco_cartao"
_DINHEIRO = re.compile(r"-?\d{1,3}(?:\.\d{3})*,\d{2}")
_DATA = re.compile(r"^\d{2}/\d{2}/\d{4}$")


def eh_extrato_pdf(texto: str) -> bool:
    return "Bradesco" in texto and re.search(r"Extrato de:\s*Ag[êe]ncia", texto) is not None


def _iso(d: str) -> str:
    dia, mes, ano = d.split("/")
    return f"{ano}-{mes}-{dia}"


def _fechar_saldos(lt: Leitura, saldos: list[float | None]) -> None:
    """Cada saldo é o anterior mais o lançamento; o primeiro dá o saldo antes do extrato."""
    anterior = None
    for mv, s in zip(lt.movimentos, saldos):
        if s is not None and anterior is not None and abs(anterior + mv.valor - s) > 0.005:
            raise ValueError(f"Extrato do Bradesco não fecha em {mv.data} ({mv.descricao}): "
                             f"{anterior} + {mv.valor} dá {round(anterior + mv.valor, 2)}, o banco diz {s}.")
        anterior = s
    primeiro = saldos[0] if saldos else None
    if lt.movimentos and primeiro is not None:
        lt.totais["saldo_inicial"] = round(primeiro - lt.movimentos[0].valor, 2)
        lt.totais["data_saldo_inicial"] = lt.movimentos[0].data


# ---------------------------------------------------------------------------
# Extrato da conta (PDF com texto)
# ---------------------------------------------------------------------------
def ler_extrato_pdf(texto: str) -> Leitura:
    lt = Leitura("extrato_bradesco_pdf")
    linhas = texto.splitlines()
    cab = next((i for i, l in enumerate(linhas) if "Crédito (R$)" in l and "Débito (R$)" in l), None)
    if cab is None:
        raise ValueError("Extrato do Bradesco sem a linha de cabeçalho (Crédito, Débito, Saldo).")
    fim_cred = linhas[cab].index("Crédito (R$)") + len("Crédito (R$)")
    fim_deb = linhas[cab].index("Débito (R$)") + len("Débito (R$)")
    saldos: list[float | None] = []
    historico, data = "", None
    for l in linhas[cab + 1:]:
        s = l.strip()
        if not s:
            continue
        if s.startswith("Total"):
            vals = [numero(v) for v in _DINHEIRO.findall(s)]
            if len(vals) >= 2:
                lt.totais["entradas"], lt.totais["saidas"] = vals[0], vals[1]
            break
        if s.startswith(("REM:", "DES:")):
            if lt.movimentos:
                lt.movimentos[-1].contraparte = re.sub(r"\s+\d{2}/\d{2}$", "", s[4:]).strip()
            continue
        valores = list(_DINHEIRO.finditer(l))
        if len(valores) < 2:
            historico = s            # linha do histórico: os valores vêm na seguinte
            continue
        primeiro = s.split()[0]
        if _DATA.match(primeiro):
            data = _iso(primeiro)
        if data is None:
            raise ValueError("Extrato do Bradesco: lançamento sem data antes dele.")
        valor_m, saldo_m = valores[-2], valores[-1]
        # Crédito ou débito pela coluna: o número fica alinhado à direita do título dela.
        credito = abs(valor_m.end() - fim_cred) <= abs(valor_m.end() - fim_deb)
        valor = numero(valor_m.group())
        lt.movimentos.append(Movimento(data, CONTA, valor if credito else -valor, historico))
        saldos.append(numero(saldo_m.group()))
        historico = ""
    _fechar_saldos(lt, saldos)
    if lt.totais.get("entradas") is not None:
        ent = round(sum(m.valor for m in lt.movimentos if m.valor > 0), 2)
        sai = round(-sum(m.valor for m in lt.movimentos if m.valor < 0), 2)
        if abs(ent - lt.totais["entradas"]) > 0.005 or abs(sai - lt.totais["saidas"]) > 0.005:
            raise ValueError(f"Extrato do Bradesco não fecha: créditos lidos {ent} x banco {lt.totais['entradas']}, "
                             f"débitos lidos {sai} x banco {lt.totais['saidas']}")
    return lt


# ---------------------------------------------------------------------------
# Extrato da conta (.csv ou .xlsx exportado pelo site/aplicativo)
# ---------------------------------------------------------------------------
def eh_extrato_tabela(linhas: list[list[str]]) -> bool:
    return bool(linhas) and str(linhas[0][0] or "").strip().startswith("Extrato de:")


def ler_extrato_tabela(linhas: list[list[str]], tipo: str) -> Leitura:
    lt = Leitura(tipo)
    cab = next((i for i, r in enumerate(linhas) if r and str(r[0]).strip() == "Data"
                and any("Histórico" in str(c) for c in r)), None)
    if cab is None:
        raise ValueError("Extrato do Bradesco sem a linha Data;Histórico;...;Saldo.")
    saldos: list[float | None] = []
    for r in linhas[cab + 1:]:
        r = [str(c if c is not None else "").strip() for c in r] + [""] * 6
        if not _DATA.match(r[0]):
            break                    # acabou a tabela (depois vêm avisos e lançamentos futuros)
        cred = numero(r[3]) if r[3] else 0.0
        deb = numero(r[4]) if r[4] else 0.0
        lt.movimentos.append(Movimento(_iso(r[0]), CONTA, round(cred - deb, 2), r[1]))
        saldos.append(numero(r[5]) if r[5] else None)
    if not lt.movimentos:
        raise ValueError("Extrato do Bradesco sem lançamentos.")
    _fechar_saldos(lt, saldos)
    return lt


def linhas_csv(texto: str) -> list[list[str]]:
    return [l.split(";") for l in texto.splitlines()]


def linhas_xlsx(conteudo: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
    if ws is None:
        return []

    def texto(c):
        if isinstance(c, (int, float)):          # o Excel guarda número: volta ao formato do extrato
            return f"{c:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        if hasattr(c, "strftime"):
            return c.strftime("%d/%m/%Y")
        return "" if c is None else str(c)
    return [[texto(c) for c in r] for r in ws.iter_rows(values_only=True)]


# ---------------------------------------------------------------------------
# Fatura do cartão em aberto (PDF de imagem, aplicativo Bradesco)
# ---------------------------------------------------------------------------
def tem_ocr() -> bool:
    return shutil.which("tesseract") is not None and shutil.which("pdftoppm") is not None


def _tesseract(img, *args: str) -> str:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    r = subprocess.run(["tesseract", "stdin", "stdout", "-l", "por", *args], input=buf.getvalue(),
                       capture_output=True, check=True)
    return r.stdout.decode("utf-8", "replace")


def _preparar(img):
    from PIL import Image
    g = img.convert("L")
    g = g.resize((g.width * 2, g.height * 2), Image.Resampling.LANCZOS)
    return g.point(lambda v: 255 if v > 170 else 0)


def _reler(pagina, caixa, permitidos: str) -> str:
    """Relê um pedaço da página (data ou valor) maior e só com os caracteres que ele pode ter."""
    x0, y0, x1, y1 = caixa
    pedaco = pagina.crop((max(0, x0 - 12), max(0, y0 - 10), x1 + 12, y1 + 10))
    pedaco = pedaco.resize((pedaco.width * 2, pedaco.height * 2))
    return _tesseract(pedaco, "--psm", "7", "-c", f"tessedit_char_whitelist={permitidos}").strip()


def _linhas_ocr(pagina) -> list[list[dict]]:
    """Palavras da página agrupadas por linha (TSV do tesseract), na ordem de leitura."""
    tsv = _tesseract(pagina, "--psm", "6", "tsv").splitlines()
    cab = tsv[0].split("\t")
    grupos: dict[tuple, list[dict]] = {}
    for l in tsv[1:]:
        c = dict(zip(cab, l.split("\t")))
        if c.get("level") != "5" or not c.get("text", "").strip():
            continue
        p = {"t": c["text"], "x": int(c["left"]), "y": int(c["top"]), "w": int(c["width"]), "h": int(c["height"])}
        grupos.setdefault((c["block_num"], c["par_num"], c["line_num"]), []).append(p)
    return sorted(grupos.values(), key=lambda ws: min(w["y"] for w in ws))


def _valor_ocr(s: str) -> float | None:
    s = s.replace("R$", "").replace(" ", "").replace("−", "-")
    s = s.translate(str.maketrans({"B": "8", "O": "0", "o": "0", "S": "5", "l": "1", "I": "1", "|": "1", "Z": "2"}))
    s = re.sub(r"[^0-9,.\-]", "", s)
    if "," not in s and re.fullmatch(r"-?\d{3,}", s.replace(".", "")):
        d = s.replace(".", "")
        s = d[:-2] + "," + d[-2:]           # o OCR às vezes come a vírgula; o real sempre tem centavos
    return numero(s) if re.fullmatch(r"-?[\d.]*\d,\d{2}", s) else None


def ler_fatura_pdf(conteudo: bytes, hoje: date | None = None) -> Leitura:
    from PIL import Image
    if not tem_ocr():
        raise ValueError("A fatura do Bradesco é imagem: precisa do tesseract (com português) e do pdftoppm instalados.")
    lt = Leitura("fatura_bradesco_pdf")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["pdftoppm", "-r", "72", "-png", "-", f"{tmp}/p"], input=conteudo, check=True, capture_output=True)
        paginas = [_preparar(Image.open(f)) for f in sorted(Path(tmp).glob("p-*.png"))]
    ref = None
    secoes: list[dict] = []
    for pagina in paginas:
        linhas = _linhas_ocr(pagina)
        alturas = [w["h"] for ws in linhas for w in ws if re.match(r"\d", w["t"])]
        normal = statistics.median(alturas) if alturas else 0
        for ws in linhas:
            texto = " ".join(w["t"] for w in ws)
            if ref is None and (m := re.search(r"Data:\s*(\d{2}/\d{2}/\d{4})", texto)):
                ref = date.fromisoformat(_iso(m[1]))
            if "Histórico" in texto and "Cota" in texto:
                # Títulos das colunas: um por cartão. O nome e o número ficam no quadro logo acima.
                y = min(w["y"] for w in ws)
                quadro = pagina.crop((int(pagina.width * 0.7), max(0, y - 4 * normal), pagina.width, y))
                quadro = quadro.resize((quadro.width * 2, quadro.height * 2))
                acima = _tesseract(quadro, "--psm", "7", "-c", "tessedit_char_whitelist=0123456789X.")
                m = re.search(r"(\d{4})\s*$", acima.strip())
                secoes.append({"final": m[1] if m else "?", "linhas": [], "total": None})
                continue
            if texto.startswith("Total para") and secoes:
                v = re.search(r"R\$\s*(-?[\d.,]+)\s*$", texto)
                secoes[-1]["total"] = _valor_ocr(v[1]) if v else 0.0
                continue
            if "USD" not in texto or not secoes or not re.match(r"\d", ws[0]["t"]):
                continue
            if min(w["y"] for w in ws) < normal / 2:
                continue             # linha cortada no alto da página: é o pedaço repetido da última da anterior
            usd = next(i for i, w in enumerate(ws) if "USD" in w["t"])
            ult_rs = max(i for i, w in enumerate(ws) if "R$" in w["t"] or "Ry" in w["t"])
            valor_ws = ws[ult_rs:]
            caixa_valor = (min(w["x"] for w in valor_ws), min(w["y"] for w in valor_ws),
                           max(w["x"] + w["w"] for w in valor_ws), max(w["y"] + w["h"] for w in valor_ws))
            d = ws[0]
            data_txt = _reler(pagina, (d["x"], d["y"], d["x"] + d["w"], d["y"] + d["h"]), "0123456789/")
            valor = _valor_ocr(_reler(pagina, caixa_valor, "0123456789,.-R$"))
            if valor is None:
                valor = _valor_ocr(" ".join(w["t"] for w in valor_ws))
            desc = " ".join(w["t"] for w in ws[1:usd]).strip(" —-–")
            desc = re.sub(r"^1(?=IOF)", "", desc)
            if not re.fullmatch(r"\d{2}/\d{2}", data_txt) or valor is None:
                raise ValueError(f"Fatura do Bradesco: não deu para ler a linha \"{texto}\" (data {data_txt!r}).")
            secoes[-1]["linhas"].append((data_txt, desc, valor))
    ref = ref or hoje or date.today()
    anterior = 0.0
    for s in secoes:
        soma = round(sum(v for _, _, v in s["linhas"]), 2)
        if s["linhas"] and (s["total"] is None or abs(soma - s["total"]) > 0.005):
            raise ValueError(f"Fatura do Bradesco não fecha no cartão final {s['final']}: lido {soma}, "
                             f"total impresso {s['total']}. Confira o arquivo (o .pdf é imagem e foi lido por OCR).")
        for d, desc, valor in s["linhas"]:
            if desc.upper().startswith("SALDO ANTERIOR"):
                anterior += valor    # dívida da fatura passada: entra no total do cartão, mas não é compra
                continue
            dia, mes = int(d[:2]), int(d[3:])
            ano = ref.year - 1 if mes > ref.month else ref.year
            # Na fatura, compra vem positiva e pagamento negativo; na conta do cartão, compra é dívida.
            lt.movimentos.append(Movimento(f"{ano}-{mes:02d}-{dia:02d}", CARTAO, -valor, desc))
    if lt.movimentos:
        # Cartão novo aqui começa devendo o saldo anterior; o pagamento antecipado dele vem nas linhas.
        lt.totais["saldo_inicial"] = round(-anterior, 2)
        lt.totais["data_saldo_inicial"] = min(m.data for m in lt.movimentos)
    return lt
