"""Exporta tudo para planilha: .xlsx simples (uma aba por tabela) ou .ods no modelo da planilha antiga.

O .ods parte do modelo original (Budget Dashboard, Tracking, Planning, Despesas Fixas) e é preenchido
pelo LibreOffice (app/ods_preencher.py, que roda com o python do sistema, o que tem o módulo uno).
"""

import io
import json
import os
import shutil
import subprocess
import tempfile
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app import db, servicos

MOEDA = '#,##0.00;[Red]-#,##0.00'
CAB = PatternFill("solid", fgColor="2F3B4C")
AMBAR = PatternFill("solid", fgColor="FFC94D")
# O modelo (Budget Dashboard da The Office Lab) é comercial e não vem no repositório: coloque o seu aqui
# e rode scripts/limpar_modelo_ods.py nele uma vez.
MODELO = Path(__file__).parent / "modelo" / "planilha.ods"
PREENCHER = Path(__file__).parent / "ods_preencher.py"
# Linhas de cada bloco do Budget Planning que o Dashboard do modelo enxerga (Income C10:C16,
# Expenses C23:C35, Savings C39:C40); o excedente vira uma linha só no fim.
CAPACIDADE = {"renda": 7, "gasto": 13, "poupanca": 2}
TIPO = {"renda": "Income", "gasto": "Expenses", "poupanca": "Savings"}


def abas(con, hoje: date | None = None) -> list[dict]:
    """As tabelas exportadas, cada uma vira uma aba. `destaque` = linhas pendentes (âmbar)."""
    hoje = hoje or date.today()
    cats = {r["id"]: r["nome"] for r in con.execute("SELECT id, nome FROM categorias")}
    pessoas = {r["id"]: r["nome"] for r in con.execute("SELECT id, nome FROM pessoas")}
    metas_nome = {r["id"]: r["nome"] for r in con.execute("SELECT id, nome FROM metas")}
    linhas, pend = [], []
    for l in servicos.lancamentos(con, limite=100000):
        for p in l["partes"] or [{"valor": l["valor"], "natureza": "", "categoria": None, "pessoa": None, "meta": None, "obs": "", "motivo": ""}]:
            linhas.append([l["data"], l["conta"], l["descricao"], l["contraparte"], p["valor"], p["natureza"],
                           cats.get(p["categoria"]) or metas_nome.get(p.get("meta")) or "", pessoas.get(p["pessoa"], ""),
                           " · ".join(x for x in (l["nota"], p["obs"]) if x), p.get("motivo", ""), l["origem"]])
            pend.append(l["pendente"])
    out = [{"nome": "Lançamentos", "cab": ["Data", "Conta", "Descrição", "Contraparte", "Valor", "Natureza", "Categoria ou meta", "Pessoa",
                                          "Obs", "Motivo", "Origem"], "linhas": linhas, "moeda": [5], "larg": [11, 14, 38, 28, 12, 13, 26, 16, 34, 30, 12],
            "destaque": [i for i, p in enumerate(pend) if p]}]

    serie = servicos.serie_mensal(con, 36)
    out.append({"nome": "Resumo mensal", "cab": ["Mês", "Renda", "Gasto", "Juros e taxas", "Resultado"],
                "linhas": [[s["mes"], s["renda"], s["gasto"], s["juros"], s["resultado"]] for s in serie], "moeda": [2, 3, 4, 5],
                "larg": [10, 12, 12, 14, 12]})
    meses = [s["mes"] for s in serie]
    tabela = defaultdict(dict)
    for m in meses:
        pi = servicos.periodo_info(m, hoje)
        for c in servicos.gastos_por_categoria(con, pi["inicio"], pi["fim"], pi["meses_orc"]):
            tabela[c["nome"]][m] = round(c["gasto"] + c["parcelas"], 2)
    out.append({"nome": "Categorias por mês", "cab": ["Categoria"] + meses + ["Total"],
                "linhas": [[nome] + [tabela[nome].get(m, 0) for m in meses] + [round(sum(tabela[nome].values()), 2)] for nome in tabela],
                "moeda": list(range(2, len(meses) + 3)), "larg": [32] + [11] * (len(meses) + 1)})
    out.append({"nome": "Pessoas", "cab": ["Pessoa", "Saldo (+ te deve, - você deve)", "Parcelas a vencer"],
                "linhas": [[p["nome"], p["saldo"], p["a_vencer"]] for p in servicos.saldo_pessoas(con)], "moeda": [2, 3], "larg": [30, 26, 18]})
    out.append({"nome": "Compromissos", "cab": ["Mês de cobrança", "Descrição", "Valor", "Parcela"],
                "linhas": [[c["mes"], c["descricao"], c["valor"], c["parcela"]] for c in servicos.compromissos(con)], "moeda": [3],
                "larg": [16, 44, 12, 10]})
    pl = servicos.plano(con, hoje)
    out.append({"nome": "Projeção", "cab": ["Mês", "Renda", "Custo de vida", "Cartão", "Reembolsos", "Entradas/saídas únicas",
                                            "Aportes em metas", "Resultado", "Falta cobrir", "Pago a pessoas", "Ainda devo", "Saldo livre",
                                            "Guardado nas metas"],
                "linhas": [[l["mes"], l["renda"], l["custo"], l["cartao"], l["reembolsos"], l["unicos"], l["aportes"], l["resultado"], l["falta"],
                            l["pago_pessoas"], l["deve_pessoas"], l["livre"], l["metas"]] for l in pl["linhas"]],
                "moeda": list(range(2, 14)), "larg": [10] + [15] * 12})
    out.append({"nome": "Eventos", "cab": ["Evento", "Tipo", "De", "Até", "Valor", "Ativo"],
                "linhas": [[e["nome"], e["tipo"], e["inicio"], e["fim"] or "", e["valor"], "sim" if e["ativo"] else "não"]
                           for e in con.execute("SELECT * FROM eventos ORDER BY inicio")], "moeda": [5], "larg": [40, 16, 10, 10, 12, 8]})
    out.append({"nome": "Metas", "cab": ["Meta", "Objetivo", "Prazo", "Aporte mensal", "Guardado", "% do objetivo", "Falta"],
                "linhas": [[m["nome"], m["alvo"], m["data_alvo"], m["aporte_mensal"], m["guardado_total"], m["completude_alvo"], m["falta_alvo"]]
                           for m in servicos.metas(con)], "moeda": [2, 4, 5, 7], "larg": [30, 12, 10, 14, 12, 14, 12]})
    return out


def planilha(con) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for a in abas(con):
        ws = wb.create_sheet(a["nome"])
        ws.append(a["cab"])
        for c in ws[1]:
            c.fill, c.font = CAB, Font(bold=True, color="FFFFFF")
        for l in a["linhas"]:
            ws.append(l)
        for col in a["moeda"]:
            for cel in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cel[0].number_format = MOEDA
        for i, w in enumerate(a["larg"]):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
        for i in a.get("destaque", []):
            ws.cell(i + 2, 7).fill = AMBAR
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Modelo antigo (.ods)
# ---------------------------------------------------------------------------
def _dobrar(itens: list[dict], cap: int, nome: str) -> list[dict]:
    """O modelo tem linhas contadas por bloco; o que passar vira uma linha só no fim."""
    if len(itens) <= cap:
        return itens
    resto, meses = itens[cap - 1:], {}
    for it in resto:
        for m, v in it["meses"].items():
            if v is not None:
                meses[m] = round((meses.get(m) or 0) + v, 2)
    return itens[:cap - 1] + [{"nome": nome, "meses": meses}]


def dados_modelo(con, hoje: date | None = None) -> dict:
    """Tudo o que o modelo precisa, já no vocabulário dele (Income/Expenses/Savings; gastos por grupo)."""
    hoje = hoje or date.today()
    prim = con.execute("SELECT MIN(data) FROM lancamentos").fetchone()[0]
    ano_ini = int(prim[:4]) if prim else hoje.year
    horizonte = db.parametro(con, "horizonte", None) or f"{hoje.year}-12"
    ano_fim = min(max(hoje.year, int(horizonte[:4])), ano_ini + 9)       # o modelo tem 10 anos
    meses = [f"{a}-{m:02d}" for a in range(ano_ini, ano_fim + 1) for m in range(1, 13)]
    orc = servicos.orcado(con, meses)
    cats = [dict(r) for r in con.execute("SELECT * FROM categorias ORDER BY ordem")]
    grupo_de = {c["id"]: c["grupo"] for c in cats}
    nome_de = {c["id"]: c["nome"] for c in cats}
    grupos = list(dict.fromkeys(c["grupo"] for c in cats if c["tipo"] == "gasto"))

    renda = [{"nome": c["nome"], "meses": orc[c["id"]]} for c in cats if c["tipo"] == "renda"]
    gastos = []
    for g in grupos:
        ids = [c["id"] for c in cats if c["tipo"] == "gasto" and c["grupo"] == g]
        vs = {m: [orc[i][m] for i in ids if orc[i][m] is not None] for m in meses}
        gastos.append({"nome": g, "meses": {m: round(sum(v), 2) if v else None for m, v in vs.items()}})
    metas = [dict(r) for r in con.execute("SELECT * FROM metas ORDER BY ativa DESC, ordem, id")]
    usadas = {r[0] for r in con.execute("SELECT DISTINCT meta FROM partes WHERE meta IS NOT NULL")}
    poupanca = [{"nome": m["nome"], "meses": orc[f"meta:{m['id']}"]} for m in metas if m["ativa"] or m["id"] in usadas]
    nome_meta = {m["id"]: m["nome"] for m in metas}

    # Budget Tracking: só o que é renda, gasto ou poupança (transferência e acerto ficam na aba Lançamentos).
    tracking = []
    for l in servicos.lancamentos(con, limite=100000):
        partes = [p for p in l["partes"] if p["natureza"] in TIPO]
        if not l["partes"]:
            tracking.append({"data": l["data"], "tipo": "Expenses" if l["valor"] < 0 else "Income", "categoria": "",
                             "valor": round(abs(l["valor"]), 2), "detalhes": f"{l['descricao']} (pendente)", "pendente": True})
        for p in partes:
            if p["natureza"] == "gasto":
                cat, det = grupo_de.get(p["categoria"], ""), f"{nome_de.get(p['categoria'], 'sem categoria')}: {l['descricao']}"
            elif p["natureza"] == "renda":
                cat, det = nome_de.get(p["categoria"], ""), l["descricao"]
            else:
                cat, det = nome_meta.get(p["meta"], ""), l["descricao"]
            sinal = 1 if p["natureza"] == "renda" else -1
            tracking.append({"data": l["data"], "tipo": TIPO[p["natureza"]], "categoria": cat, "valor": round(sinal * p["valor"], 2),
                             "detalhes": det + "".join(f" ({x})" for x in (l["nota"], p["obs"]) if x), "pendente": l["pendente"]})

    # Despesas Fixas: um bloco por mês do ano corrente, itens agrupados como as linhas do Budget Planning.
    ano = [f"{hoje.year}-{m:02d}" for m in range(1, 13)]
    itens = servicos.itens_fixos(con)
    fixas = []
    for g in grupos:
        its = [i for c in cats if c["tipo"] == "gasto" and c["grupo"] == g for i in itens[c["id"]]]
        fixas.append({"nome": g, "itens": [{"nome": i["nome"], "valores": {m: i["valor"] for m in ano
                                                                           if servicos._no_intervalo(m, i["inicio"] or "", i["fim"])}}
                                          for i in its]})
    return {
        "ano_inicial": ano_ini, "ano": hoje.year, "mes": hoje.month,
        "dia_corte": int(db.parametro(con, "dia_renda_mes_seguinte", "25")),
        "renda": _dobrar(renda, CAPACIDADE["renda"], "Outras entradas"),
        "gastos": _dobrar(gastos, CAPACIDADE["gasto"], "Outros grupos"),
        "poupanca": _dobrar(poupanca, CAPACIDADE["poupanca"], "Outras metas") or [{"nome": "", "meses": {}}],
        "tracking": tracking, "fixas": {"meses": ano, "grupos": fixas}, "abas": abas(con, hoje),
    }


def planilha_modelo(dados: dict) -> bytes:
    """Roda o LibreOffice sobre o modelo com os dados de dados_modelo() (uns 15 s)."""
    if not MODELO.exists():
        raise RuntimeError("falta o modelo em app/modelo/planilha.ods (ele não vem no repositório; veja o README)")
    with tempfile.TemporaryDirectory(prefix="financas-ods-") as tmp:
        entrada, modelo, saida = Path(tmp) / "dados.json", Path(tmp) / "modelo.ods", Path(tmp) / "planilha.ods"
        entrada.write_text(json.dumps(dados, ensure_ascii=False))
        # Cópia gravável: o modelo no container é do root, e o LibreOffice abre arquivo sem permissão de
        # escrita como somente leitura, recusando qualquer alteração.
        shutil.copyfile(MODELO, modelo)
        env = {**os.environ, "HOME": tmp, "LANG": "pt_BR.UTF-8", "LC_ALL": "pt_BR.UTF-8"}
        r = subprocess.run(["/usr/bin/python3", str(PREENCHER), str(entrada), str(modelo), str(saida), tmp],
                           capture_output=True, text=True, timeout=240, env=env)
        if r.returncode != 0 or not saida.exists():
            raise RuntimeError("LibreOffice não gerou a planilha: " + (r.stderr or r.stdout)[-800:])
        return saida.read_bytes()
